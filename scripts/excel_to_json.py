import json
import sys

from openpyxl import load_workbook


def cell_value(value):
    if value is None:
        return ""
    return str(value).strip()


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: excel_to_json.py workbook.xlsx")

    workbook = load_workbook(sys.argv[1], read_only=True, data_only=True)
    requested_sheet = sys.argv[2] if len(sys.argv) > 2 else workbook.sheetnames[0]
    if requested_sheet not in workbook.sheetnames:
        print("[]")
        return
    sheet = workbook[requested_sheet]
    rows = []
    for row in sheet.iter_rows(values_only=True):
        values = [cell_value(value) for value in row]
        if any(values):
            rows.append(values)
    print(json.dumps(rows, ensure_ascii=False))


if __name__ == "__main__":
    main()
